from pathlib import Path

from docx import Document as WordDocument
import pymupdf
import pytest

from app.document_processing.parsers.docx import parse_docx
from app.document_processing.parsers.markdown import parse_markdown
from app.document_processing.parsers.pdf import parse_pdf
from app.document_processing.parsers.text import parse_text
from app.document_processing.schemas import ParsedTextUnit


def test_pdf_parser_extracts_one_unit_per_page(tmp_path: Path) -> None:
    pdf_path = tmp_path / "lecture.pdf"
    document = pymupdf.open()
    first_page = document.new_page()
    first_page.insert_text((72, 72), "First page")
    second_page = document.new_page()
    second_page.insert_text((72, 72), "Second page")
    document.save(pdf_path)
    document.close()

    units = parse_pdf(pdf_path)

    assert [unit.text.strip() for unit in units] == ["First page", "Second page"]
    assert [unit.sequence for unit in units] == [0, 1]
    assert [unit.source_type for unit in units] == ["page", "page"]
    assert [(unit.source_start, unit.source_end) for unit in units] == [
        (1, 1),
        (2, 2),
    ]


def test_docx_parser_extracts_body_paragraphs_in_order(tmp_path: Path) -> None:
    docx_path = tmp_path / "lecture.docx"
    document = WordDocument()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(docx_path)

    units = parse_docx(docx_path)

    assert units == [
        ParsedTextUnit(
            text="First paragraph",
            sequence=0,
            source_type="paragraph",
            source_start=1,
            source_end=1,
        ),
        ParsedTextUnit(
            text="Second paragraph",
            sequence=1,
            source_type="paragraph",
            source_start=2,
            source_end=2,
        ),
    ]


def test_text_parser_accepts_utf_8_bom_and_retains_line_numbers(
    tmp_path: Path,
) -> None:
    text_path = tmp_path / "notes.txt"
    text_path.write_bytes(b"\xef\xbb\xbfFirst line\nSecond line\r\n")

    units = parse_text(text_path)

    assert units == [
        ParsedTextUnit(
            text="First line",
            sequence=0,
            source_type="line",
            source_start=1,
            source_end=1,
        ),
        ParsedTextUnit(
            text="Second line",
            sequence=1,
            source_type="line",
            source_start=2,
            source_end=2,
        ),
    ]


@pytest.mark.parametrize(
    "original_text",
    [
        pytest.param("# Heading\n\n- **important** item\n", id="lf"),
        pytest.param("# Heading\r\n\r\n- **important** item\r\n", id="crlf"),
    ],
)
def test_markdown_parser_preserves_original_text_and_line_range(
    tmp_path: Path,
    original_text: str,
) -> None:
    markdown_path = tmp_path / "notes.md"
    markdown_path.write_bytes(original_text.encode("utf-8"))

    units = parse_markdown(markdown_path)

    assert units == [
        ParsedTextUnit(
            text=original_text,
            sequence=0,
            source_type="line_range",
            source_start=1,
            source_end=3,
        )
    ]
